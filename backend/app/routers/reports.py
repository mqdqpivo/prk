import io
from datetime import date
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func
from typing import Optional

from app.database import get_db
from app.models import (
    User, UserRole, Grade, Attendance, Student, Subject, Group,
    TeacherSubjectGroup, Teacher
)
from app.dependencies import get_current_user, require_roles

router = APIRouter(prefix="/api/reports", tags=["reports"])


def _generate_gradebook_excel(db, group_id, subject_id, semester):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал оценок"

    group = db.query(Group).filter(Group.id == group_id).first()
    subject = db.query(Subject).filter(Subject.id == subject_id).first()

    header_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    header_font = Font(bold=True)

    ws.append([f"Журнал успеваемости | {group.name if group else ''} | {subject.name if subject else ''} | Семестр {semester or ''}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)

    headers = ["ФИО", "Дата", "Тип", "Оценка", "Комментарий"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    q = db.query(Grade).options(
        joinedload(Grade.student).joinedload(Student.user),
        joinedload(Grade.subject)
    ).filter(Grade.subject_id == subject_id)

    if group_id:
        q = q.join(Student).filter(Student.group_id == group_id)

    grade_colors = {
        5: "DCFCE7", 4: "FEF9C3", 3: "FFEDD5", 2: "FEE2E2"
    }

    for grade in q.order_by(Grade.grade_date).all():
        user = grade.student.user if grade.student else None
        name = f"{user.last_name} {user.first_name} {user.middle_name or ''}" if user else ""
        row_idx = ws.max_row + 1
        ws.append([
            name.strip(),
            str(grade.grade_date) if grade.grade_date else "",
            grade.grade_type.value if grade.grade_type else "",
            float(grade.value) if grade.value else "",
            grade.comment or ""
        ])
        val = int(float(grade.value)) if grade.value else 0
        if val in grade_colors:
            ws.cell(row=row_idx, column=4).fill = PatternFill(
                start_color=grade_colors[val], end_color=grade_colors[val], fill_type="solid"
            )

    ws.freeze_panes = "A3"
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _generate_gradebook_pdf(db, group_id, subject_id, semester):
    from weasyprint import HTML

    group = db.query(Group).filter(Group.id == group_id).first()
    subject = db.query(Subject).filter(Subject.id == subject_id).first()

    q = db.query(Grade).options(
        joinedload(Grade.student).joinedload(Student.user),
    ).filter(Grade.subject_id == subject_id)
    if group_id:
        q = q.join(Student).filter(Student.group_id == group_id)

    rows = ""
    for g in q.order_by(Grade.grade_date).all():
        user = g.student.user if g.student else None
        name = f"{user.last_name} {user.first_name}" if user else ""
        rows += f"<tr><td>{name}</td><td>{g.grade_date or ''}</td><td>{g.grade_type.value if g.grade_type else ''}</td><td>{float(g.value) if g.value else ''}</td><td>{g.comment or ''}</td></tr>"

    html_str = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>
body {{ font-family: sans-serif; font-size: 12px; }}
h1 {{ font-size: 16px; text-align: center; }}
table {{ width: 100%; border-collapse: collapse; }}
th, td {{ border: 1px solid #ccc; padding: 4px 8px; text-align: left; }}
th {{ background: #DBEAFE; }}
</style></head><body>
<h1>ИМСИТ | Журнал успеваемости | {group.name if group else ''} | {subject.name if subject else ''} | Семестр {semester or ''}</h1>
<table><tr><th>ФИО</th><th>Дата</th><th>Тип</th><th>Оценка</th><th>Комментарий</th></tr>{rows}</table>
</body></html>"""

    buf = io.BytesIO()
    HTML(string=html_str).write_pdf(buf)
    buf.seek(0)
    return buf


@router.get("/gradebook")
def report_gradebook(
    group_id: int = Query(...),
    subject_id: int = Query(...),
    semester: Optional[int] = None,
    format: str = Query("xlsx"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if format == "pdf":
        buf = _generate_gradebook_pdf(db, group_id, subject_id, semester)
        return StreamingResponse(buf, media_type="application/pdf",
                                 headers={"Content-Disposition": "attachment; filename=gradebook.pdf"})
    else:
        buf = _generate_gradebook_excel(db, group_id, subject_id, semester)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=gradebook.xlsx"})


@router.get("/attendance")
def report_attendance(
    group_id: int = Query(...),
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    format: str = Query("xlsx"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    group = db.query(Group).filter(Group.id == group_id).first()

    q = db.query(Attendance).options(
        joinedload(Attendance.student).joinedload(Student.user),
        joinedload(Attendance.schedule_slot)
    ).join(Student).filter(Student.group_id == group_id)
    if date_from:
        q = q.filter(Attendance.lesson_date >= date_from)
    if date_to:
        q = q.filter(Attendance.lesson_date <= date_to)

    if format == "pdf":
        from weasyprint import HTML
        rows = ""
        for a in q.order_by(Attendance.lesson_date).all():
            user = a.student.user if a.student else None
            name = f"{user.last_name} {user.first_name}" if user else ""
            status_map = {"present": "П", "absent": "О", "late": "Оп", "excused": "Ув"}
            st = status_map.get(a.status.value if a.status else "", "")
            rows += f"<tr><td>{name}</td><td>{a.lesson_date or ''}</td><td>{st}</td></tr>"

        html_str = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>
body {{ font-family: sans-serif; font-size: 12px; }}
h1 {{ font-size: 16px; text-align: center; }}
table {{ width: 100%; border-collapse: collapse; }}
th, td {{ border: 1px solid #ccc; padding: 4px 8px; }}
th {{ background: #DBEAFE; }}
</style></head><body>
<h1>ИМСИТ | Журнал посещаемости | {group.name if group else ''}</h1>
<table><tr><th>ФИО</th><th>Дата</th><th>Статус</th></tr>{rows}</table>
</body></html>"""
        buf = io.BytesIO()
        HTML(string=html_str).write_pdf(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf",
                                 headers={"Content-Disposition": "attachment; filename=attendance.pdf"})

    wb = Workbook()
    ws = wb.active
    ws.title = "Посещаемость"
    header_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

    ws.append([f"Журнал посещаемости | {group.name if group else ''}"])
    headers = ["ФИО", "Дата", "Статус"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=2, column=col_idx).fill = header_fill
        ws.cell(row=2, column=col_idx).font = Font(bold=True)

    status_map = {"present": "П", "absent": "О", "late": "Оп", "excused": "Ув"}
    for a in q.order_by(Attendance.lesson_date).all():
        user = a.student.user if a.student else None
        name = f"{user.last_name} {user.first_name}" if user else ""
        st = status_map.get(a.status.value if a.status else "", "")
        ws.append([name, str(a.lesson_date or ""), st])

    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=attendance.xlsx"})


@router.get("/student/{student_id}")
def report_student(
    student_id: int,
    format: str = Query("xlsx"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    student = db.query(Student).options(
        joinedload(Student.user), joinedload(Student.group).joinedload(Group.specialty)
    ).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Студент не найден")

    subjects_data = db.query(
        Subject.name,
        func.count(Grade.id).label("grade_count"),
        func.avg(Grade.value).label("avg_grade"),
    ).join(Grade, Grade.subject_id == Subject.id).filter(
        Grade.student_id == student_id
    ).group_by(Subject.name).all()

    attendance_data = db.query(
        Subject.name,
        func.count(Attendance.id).label("absent_count"),
    ).select_from(Attendance).join(
        Student, Attendance.student_id == Student.id
    ).filter(
        Attendance.student_id == student_id,
        Attendance.status == "absent"
    ).group_by(Subject.name).all() if False else []

    user = student.user
    group = student.group

    if format == "pdf":
        from weasyprint import HTML
        rows = ""
        for s in subjects_data:
            rows += f"<tr><td>{s.name}</td><td>{s.grade_count}</td><td>{round(float(s.avg_grade), 2) if s.avg_grade else '-'}</td></tr>"

        html_str = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><style>
body {{ font-family: sans-serif; font-size: 12px; }}
h1 {{ font-size: 16px; text-align: center; }}
table {{ width: 100%; border-collapse: collapse; }}
th, td {{ border: 1px solid #ccc; padding: 4px 8px; }}
th {{ background: #DBEAFE; }}
.info {{ margin-bottom: 20px; }}
</style></head><body>
<h1>ИМСИТ | Индивидуальный отчёт студента</h1>
<div class="info">
<p><b>ФИО:</b> {user.last_name} {user.first_name} {user.middle_name or ''}</p>
<p><b>Группа:</b> {group.name if group else ''}</p>
<p><b>Номер студ. билета:</b> {student.student_number or ''}</p>
</div>
<table><tr><th>Предмет</th><th>Оценок</th><th>Средний балл</th></tr>{rows}</table>
</body></html>"""
        buf = io.BytesIO()
        HTML(string=html_str).write_pdf(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf",
                                 headers={"Content-Disposition": f"attachment; filename=student_{student_id}.pdf"})

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт студента"
    header_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

    ws.append([f"{user.last_name} {user.first_name} {user.middle_name or ''} | {group.name if group else ''} | {student.student_number or ''}"])
    ws.append(["Предмет", "Оценок", "Средний балл"])
    for col_idx in range(1, 4):
        ws.cell(row=2, column=col_idx).fill = header_fill
        ws.cell(row=2, column=col_idx).font = Font(bold=True)

    for s in subjects_data:
        ws.append([s.name, s.grade_count, round(float(s.avg_grade), 2) if s.avg_grade else ""])

    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=student_{student_id}.xlsx"})
